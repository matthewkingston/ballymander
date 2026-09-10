#!/usr/bin/env python3
"""Parse EONI 2023 local council election result sheets into one JSON.

Extracts party + first-preference votes by DEA. Source layouts:
  A. per-DEA PDF, "Elected No Candidate Description 1st Preference"  (30 files)
  B. per-DEA PDF, "No. Candidate Party First Preference Votes"       (12 files)
  C. per-DEA xlsx                                                    (14 files)
  D. per-council xlsx, one worksheet per DEA                         (3 files)

Every DEA is reconciled against the sheet's own stated Total Valid Votes.

PDF hazards handled:
  * long party names overflow their column and interleave with the vote number
    when read left-to-right, so cells are recovered in PDF draw order instead;
  * stray overprinted glyphs ('SDLP' drawn with an 'e' over the 'S');
  * long candidate names wrapped across two physical lines;
  * the number/name columns run together ('4Kennedy, Bill').
"""
import json, re, os, glob, unicodedata
from collections import defaultdict

RAW = "/home/matt/vote/data/council_elections_23/raw"
OUT = "/home/matt/vote/data/council_elections_23/council_elections_2023.json"

class ScannedPDF(Exception):
    pass

NUM = re.compile(r"-?[\d,]+\.\d{2}|-?[\d,]+")  # 2dp values; never span two cells

def norm_ws(s):
    s = unicodedata.normalize("NFC", s or "")
    s = s.replace("–", "-").replace("—", "-").replace("’", "'")
    return re.sub(r"\s+", " ", s).strip()

def to_int(tok):
    return int(round(float(tok.replace(",", ""))))

# ------------------------------------------------------------------ PDF ----
def column_rules(page, header_top):
    """Vertical rules of the table body (the header row merges its cells)."""
    edges = []
    for r in page.rects:
        if r["bottom"] > header_top - 2:
            edges += [r["x0"], r["x1"]]
    if not edges:
        return []
    edges.sort()
    merged = [edges[0]]
    for e in edges[1:]:
        if e - merged[-1] > 2.0:
            merged.append(e)
    return merged

def row_bands(page, header_top):
    """Horizontal rules of the table body, giving each candidate row's extent.

    Rows must be recovered as bands, not as a single text baseline: a wrapped
    candidate name occupies two baselines, and in the Mid Ulster template the
    vote figures sit at a different vertical offset from the name beside them.
    """
    ys = []
    for r in page.rects:
        if r["bottom"] > header_top - 2:
            ys += [r["top"], r["bottom"]]
    ys = sorted(set(round(y, 1) for y in ys))
    if not ys:
        return []
    merged = [ys[0]]
    for y in ys[1:]:
        if y - merged[-1] > 2.0:
            merged.append(y)
    return merged

def row_cells(chars, rules):
    """Split one physical row into cells.

    Cells cannot be cut on x alone: EONI sheets let a long candidate name spill
    into the party column and a long party name spill across the vote number, so
    reading left-to-right interleaves them ('Democratic Unionist Party - D.U9.1P3..00').
    Glyphs drawn contiguously always belong to the same cell, so the column is
    only re-derived from x when there is a gap or the pen jumps backwards.
    """
    def col_of(x):
        i = 0
        for j, r in enumerate(rules):
            if x >= r - 0.5:
                i = j
        return i
    cells, prev, col = defaultdict(list), None, None
    for c in chars:
        if prev is None or c["x0"] < prev["x1"] - 1.0 or c["x0"] > prev["x1"] + 1.0:
            col = col_of(c["x0"])
        cells[col].append(c)
        prev = c
    return cells

def cell_text(cs):
    return norm_ws("".join(c["text"] for c in cs))

SKIP_ROW = re.compile(r"^\s*(non[- ]?transferable|totals?\b|totals correct|"
                      r"enter timings|result\b|stage\b)", re.I)

def parse_pdf(path):
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        page = pdf.pages[0]
        text = page.extract_text() or ""
        if len(text.strip()) < 200:
            raise ScannedPDF("no text layer (scanned image)")
        words = page.extract_words(x_tolerance=1)

        def grab(pat, cast=int):
            m = re.search(pat, text, re.I)
            return cast(m.group(1).replace(",", "")) if m else None
        meta = {
            "eligible_electorate": grab(r"Eligible Electorate\s+([\d,]+)"),
            "votes_polled":        grab(r"Votes Polled\s+([\d,]+)"),
            "valid_votes":         grab(r"Total Valid Votes\s+([\d,]+)"),
            "invalid_votes":       grab(r"Invalid Votes\s+([\d,]+)"),
            "seats":               grab(r"Number to be [Ee]lected\s+([\d,]+)"),
            "quota":               grab(r"(?:Electoral )?Quota(?: of)?\s+([\d,]+)"),
        }
        m = re.search(r"District Electoral Area of\s+(.+?)(?:\s+BACK to|\s+Stage\b|$)",
                      text, re.I | re.M)
        lines = [l for l in text.splitlines() if l.strip()]
        dea_in_file = norm_ws(m.group(1)) if m else (norm_ws(lines[0]) if lines else "")

        cand_w = [w for w in words if w["text"].rstrip(".") == "Candidate"]
        desc_w = [w for w in words if w["text"].rstrip(".") in ("Description", "Party")]
        if not cand_w or not desc_w:
            raise ValueError("no Candidate/Description|Party header")
        hdr_top = cand_w[0]["top"]
        same = [w for w in desc_w if abs(w["top"] - hdr_top) < 3]
        desc_w = same or desc_w
        rules = column_rules(page, hdr_top)
        if len(rules) < 4:
            raise ValueError("too few column rules")

        def col_of(x):
            i = 0
            for j, r in enumerate(rules):
                if x >= r - 0.5:
                    i = j
            return i
        cand_c, desc_c = col_of(cand_w[0]["x0"]), col_of(desc_w[0]["x0"])
        if desc_c <= cand_c:
            desc_c = cand_c + 1

        bands = row_bands(page, hdr_top)
        rows = defaultdict(list)
        for c in page.chars:
            mid = (c["top"] + c["bottom"]) / 2.0
            if mid <= hdr_top + 3:
                continue
            bi = None
            for j, y in enumerate(bands):
                if mid >= y - 0.5:
                    bi = j
            rows[bi if bi is not None else round(mid, 0)].append(c)

        anchors, orphans = [], []
        for key in sorted(rows, key=lambda k: (k is None, k)):
            cells = row_cells(rows[key], rules)
            name_cs, party_cs = cells.get(cand_c, []), cells.get(desc_c, [])
            # a name flush against the party column glues to it; if the party
            # cell is then empty, cut the merged run at the column rule instead
            if name_cs and not party_cs and desc_c < len(rules):
                edge = rules[desc_c]
                keep = [c for c in name_cs if c["x0"] < edge - 0.5]
                spill = [c for c in name_cs if c["x0"] >= edge - 0.5]
                if keep and spill:
                    name_cs, party_cs = keep, spill
            name, party = cell_text(name_cs), cell_text(party_cs)
            # labels such as "Non-Transferable" wrap across the name/party
            # columns, so test the joined text as well as each cell
            joined = norm_ws(name + party)
            if (SKIP_ROW.match(name) or SKIP_ROW.match(party) or SKIP_ROW.match(joined)
                    or name.lower() == "candidate"
                    or party.lower() in ("description", "party")):
                continue
            fp = None
            for ci in sorted(k for k in cells if isinstance(k, int) and k > desc_c):
                mnum = NUM.match(cell_text(cells[ci]))
                if mnum:
                    fp = to_int(mnum.group(0)); break
            name = re.sub(r"^(elected|excluded|e)\b[\s.]*", "", name, flags=re.I)
            name = re.sub(r"^\d+\s*", "", name).strip()
            top = min(c["top"] for c in rows[key])
            if party and fp is not None:
                anchors.append({"top": top, "candidate": name,
                                "party_raw": party, "first_pref": fp})
            elif name and not party:
                orphans.append({"top": top, "text": name})

        # stitch candidate names wrapped across two physical lines
        for o in orphans:
            if not anchors:
                continue
            a = min(anchors, key=lambda a: abs(a["top"] - o["top"]))
            if abs(a["top"] - o["top"]) < 26:
                a["candidate"] = norm_ws(
                    (a["candidate"] + " " + o["text"]) if a["top"] < o["top"]
                    else (o["text"] + " " + a["candidate"]))
        for a in anchors:
            a.pop("top", None)
        return dea_in_file, meta, anchors

# ----------------------------------------------------------------- XLSX ----
XL_LABELS = [
    ("eligible_electorate", r"eligible electorate"),
    ("votes_polled",        r"(total )?votes polled"),
    ("valid_votes",         r"(total )?valid votes"),
    ("invalid_votes",       r"invalid votes"),
    ("seats",               r"number (to be|of members to be) elected"),
    ("quota",               r"electoral quota"),
]

def _cells(ws):
    return [[c for c in row] for row in ws.iter_rows(values_only=True)]

def _txt(v):
    return norm_ws("" if v is None else str(v))

def parse_xlsx_sheet(ws):
    grid = _cells(ws)
    dea = None
    meta = {k: None for k, _ in XL_LABELS}
    for row in grid:
        for j, v in enumerate(row):
            t = _txt(v).lower().rstrip(": ")
            if not t:
                continue
            nxt = next((row[k] for k in range(j + 1, len(row))
                        if _txt(row[k]) != ""), None)
            if re.fullmatch(r"district electoral area( of)?", t) and dea is None:
                dea = _txt(nxt)
            for key, pat in XL_LABELS:
                if meta[key] is None and re.fullmatch(pat, t):
                    try:
                        meta[key] = int(round(float(str(nxt).replace(",", ""))))
                    except (TypeError, ValueError):
                        pass

    # header row: the one naming the candidate column
    hdr = name_c = desc_c = fp_c = None
    for i, row in enumerate(grid):
        for j, v in enumerate(row):
            t = _txt(v).lower()
            if re.search(r"^(candidate|names of candidates)", t):
                hdr, name_c = i, j
                break
        if hdr is not None:
            break
    if hdr is None:
        return None
    for j, v in enumerate(grid[hdr]):
        t = _txt(v).lower()
        if desc_c is None and re.search(r"description", t):
            desc_c = j
        if fp_c is None and re.search(r"(1st|first) preference", t):
            fp_c = j
    if fp_c is None:                      # bundle layout: sub-header row below
        for i in range(hdr + 1, min(hdr + 4, len(grid))):
            for j, v in enumerate(grid[i]):
                if re.search(r"(1st|first) preference", _txt(v).lower()):
                    fp_c = j; break
            if fp_c is not None:
                break
    if desc_c is None or fp_c is None:
        return None

    # some workbooks print the whole candidate table twice in one sheet
    stop = len(grid)
    for i in range(hdr + 1, len(grid)):
        if name_c < len(grid[i]) and _txt(grid[i][name_c]).lower() in (
                "candidate", "names of candidates"):
            stop = i; break

    cands, blanks = [], 0
    for row in grid[hdr + 1:stop]:
        name = _txt(row[name_c]) if name_c < len(row) else ""
        party = _txt(row[desc_c]) if desc_c < len(row) else ""
        raw = row[fp_c] if fp_c < len(row) else None
        if not name and not party:
            blanks += 1
            if blanks > 25 and cands:
                break
            continue
        blanks = 0
        if SKIP_ROW.match(name) or SKIP_ROW.match(party):
            continue
        try:
            fp = int(round(float(str(raw).replace(",", ""))))
        except (TypeError, ValueError):
            continue
        if name and party:
            cands.append({"candidate": name, "party_raw": party, "first_pref": fp})
    return dea, meta, cands

def parse_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    for sn in wb.sheetnames:
        got = parse_xlsx_sheet(wb[sn])
        if got and got[2]:
            out.append(got)
    wb.close()
    return out

# ------------------------------------------------------- SCANNED SHEETS ----
# The seven Fermanagh & Omagh DEAs are published as image-only PDFs with no
# text layer, so they cannot be parsed. These figures were transcribed from the
# rendered page images; each DEA's first preferences sum exactly to the Total
# Valid Votes printed on the same sheet, which is what validates them.
# Party strings are as displayed; the source renders long descriptions clipped
# (e.g. "SDLP (Social Democ"), so canonical names are recorded here.
SCANNED = {
 "Enniskillen": dict(meta=dict(eligible_electorate=13642, votes_polled=7775,
      valid_votes=7676, invalid_votes=99, seats=6, quota=1097), rows=[
   ("Blake, Paul","SDLP (Social Democratic and Labour Party)",539),
   ("Browne, Dermot","Sinn Féin",846),
   ("Crawford, Donald","TUV - No Sea Border",509),
   ("Crawford, Roy","Ulster Unionist Party",667),
   ("Elliott, Keith","Democratic Unionist Party - D.U.P.",1029),
   ("Irvine, Robert","Ulster Unionist Party",562),
   ("Mahon, Jill","Democratic Unionist Party - D.U.P.",437),
   ("Maguire, Tommy","Sinn Féin",1024),
   ("McManus, Andrea","Sinn Féin",780),
   ("O'Cofaigh, Donal","Cross-Community Labour Alternative",504),
   ("Roofe, Eddie","Alliance",779)]),
 "Erne East": dict(meta=dict(eligible_electorate=12127, votes_polled=8588,
      valid_votes=8478, invalid_votes=110, seats=6, quota=None), rows=[
   ("Bullick, Richard","Alliance Party",221),
   ("Greene, Sheamus","Sinn Féin",1753),
   ("Hayes, Noeleen","Sinn Féin",1296),
   ("Keenan, Eamon","Independent",560),
   ("McDermott, Tina","Independent",238),
   ("McPhillips, Garbhan","SDLP (Social Democratic and Labour Party)",498),
   ("O'Reilly Thomas","Sinn Féin",1387),
   ("Robinson, Paul","Democratic Unionist Party - D.U.P.",1410),
   ("Warrington, Victor","Ulster Unionist Party",1115)]),
 "Erne North": dict(meta=dict(eligible_electorate=11509, votes_polled=7193,
      valid_votes=7100, invalid_votes=93, seats=5, quota=None), rows=[
   ("Armstrong, Diana","Ulster Unionist Party",1392),
   ("Bullick, Eric","Alliance Party",358),
   ("Coyle, Debbie","Sinn Féin",1561),
   ("Coyle, John","SDLP (Social Democratic and Labour Party)",783),
   ("Elliott, Alex","TUV - No Sea Border",422),
   ("Feely, John","Sinn Féin",852),
   ("Mahon, David","Democratic Unionist Party - D.U.P.",984),
   ("McClaughry, John","Ulster Unionist Party",529),
   ("Stevenson, Paul","Democratic Unionist Party - D.U.P.",219)]),
 "Erne West": dict(meta=dict(eligible_electorate=11196, votes_polled=7652,
      valid_votes=7586, invalid_votes=66, seats=5, quota=None), rows=[
   ("Brough, Elaine","Sinn Féin",1362),
   ("Elliott, Aaron","Democratic Unionist Party - D.U.P.",687),
   ("Feely, Anthony","Sinn Féin",1493),
   ("Gannon, Adam","SDLP (Social Democratic and Labour Party)",835),
   ("McArdle, Declan","Sinn Féin",1256),
   ("McCusker, Gerard","Alliance Party",250),
   ("McGoldrick, Paul","Independent",506),
   ("Ovens, Mark","Ulster Unionist Party",1197)]),
 "Mid Tyrone": dict(meta=dict(eligible_electorate=13083, votes_polled=8661,
      valid_votes=8542, invalid_votes=119, seats=6, quota=1221), rows=[
   ("Barton, Rosemary","Ulster Unionist Party",687),
   ("Beaumont, Matthew","Alliance Party",311),
   ("Devine Gallagher, Ro","Sinn Féin",1464),
   ("Fitzgerald, Anne-Marie","Sinn Féin",1204),
   ("Hawkes, Shirley","Democratic Unionist Party - D.U.P.",1057),
   ("Kelly, Padraigin","Sinn Féin",1621),
   ("McAleer, Emmet","Independent",629),
   ("McGrath, Bernard","SDLP (Social Democratic and Labour Party)",497),
   ("Withers, Patrick","Sinn Féin",1072)]),
 "Omagh": dict(meta=dict(eligible_electorate=13279, votes_polled=6739,
      valid_votes=6636, invalid_votes=103, seats=6, quota=None), rows=[
   ("Bell, Matthew","Ulster Unionist Party",489),
   ("Deehan, Josephine","Independent",469),
   ("Donnelly, Stephen","Alliance Party",872),
   ("Dunphy, Kathy","Independent",80),
   ("Ferguson, Amy","Socialist Party (Northern Ireland)",75),
   ("Kelly, Catherine","Sinn Féin",1013),
   ("McColgan, Marty","Sinn Féin",773),
   ("McElduff, Barry","Sinn Féin",1342),
   ("Mellon, Brenda","SDLP (Social Democratic and Labour Party)",385),
   ("Thompson, Errol","Democratic Unionist Party - D.U.P.",1138)]),
 "West Tyrone": dict(meta=dict(eligible_electorate=12652, votes_polled=8413,
      valid_votes=8301, invalid_votes=112, seats=6, quota=None), rows=[
   ("Buchanan, Mark","Democratic Unionist Party - D.U.P.",1619),
   ("Campbell, Glenn","Sinn Féin",1279),
   ("Donnelly, Ann-Marie","Sinn Féin",1129),
   ("Donnelly, Joyce","Alliance Party",457),
   ("Garrity, Mary","SDLP (Social Democratic and Labour Party)",746),
   ("McCann, Stephen","Sinn Féin",1093),
   ("McNulty, Colette","Sinn Féin",1149),
   ("Rainey, Allan","Ulster Unionist Party",829)]),
}

# ------------------------------------------------------ PARTY STANDARD -----
# Party labels deliberately match data/pc24_results_2024.json so the two
# elections can be joined on `party` directly. Codes there: Sinn Féin, DUP,
# Alliance, UUP, SDLP, TUV, Independent, Green, PBP, Aontú, CCLA, Conservative.
# The last four below stood at council level only and follow the same style
# (established abbreviation where that is the common usage, else the name).
PARTY_CANON = {
  "Sinn Féin":      "Sinn Féin",
  "DUP":            "Democratic Unionist Party",
  "Alliance":       "Alliance Party of Northern Ireland",
  "UUP":            "Ulster Unionist Party",
  "SDLP":           "Social Democratic and Labour Party",
  "TUV":            "Traditional Unionist Voice",
  "Independent":    "Independent",
  "Green":          "Green Party Northern Ireland",
  "PBP":            "People Before Profit",
  "Aontú":          "Aontú",
  "CCLA":           "Cross-Community Labour Alternative",
  "Conservative":   "Conservative and Unionist Party",
  "Workers Party":  "The Workers Party",
  "PUP":            "Progressive Unionist Party",
  "IRSP":           "Irish Republican Socialist Party",
  "Socialist Party":"Socialist Party",
}
# every raw description string observed across all 80 DEAs, mapped explicitly
PARTY_MAP = {
  "dup":"DUP", "d.u.p.":"DUP", "democratic unionist party - d.u.p.":"DUP",
  "democratic unionist party - d.u.p":"DUP", "democratic unionist part - d.u.p":"DUP",
  "democratic ulster unionist - d.u.p.":"DUP", "democratic unionist party":"DUP",
  "sinn féin":"Sinn Féin", "sinn fein":"Sinn Féin", "sf":"Sinn Féin",
  "uup":"UUP", "ulster unionist party":"UUP", "ulster unionist":"UUP",
  "alliance party":"Alliance", "alliance":"Alliance", "all":"Alliance", "ap":"Alliance",
  "sdlp":"SDLP", "sdlp (social democratic & labour party)":"SDLP",
  "sdlp (social democratic and labour party)":"SDLP",
  "tuv":"TUV", "tuv - no sea border":"TUV",
  "green party northern ireland":"Green", "green party ni":"Green",
  "green party":"Green", "grn":"Green", "green part northern ireland":"Green",
  "people before profit alliance":"PBP", "pbpa":"PBP",
  "aontú":"Aontú", "aontu":"Aontú",
  "aontú for life, unity, economic justice":"Aontú",
  "the workers party":"Workers Party",
  "progressive unionist party of northern ireland":"PUP", "pup":"PUP",
  "conservative and unionist party":"Conservative", "con":"Conservative",
  "irsp - for a socialist republic":"IRSP",
  "ccla":"CCLA", "cross-community labour alternative":"CCLA",
  "socialist party (northenr ireland)":"Socialist Party",   # typo in source
  "socialist party (northern ireland)":"Socialist Party",
  "independent":"Independent", "ind":"Independent", "indepdendent":"Independent",  # typo in source
}

# --- voting modelling assumptions ------------------------------------------
# Independents counted as the party their voters would strongly tend to back at
# Westminster level, since this data exists to be extrapolated upward. These are
# deliberate modelling choices, not data corrections: the sheet's own wording is
# kept in `party_raw`, the declared label in `party_as_declared`, and the
# candidate is marked party_source="assumed". Independents not listed here are
# left as Independent by explicit decision. See "Voting modelling assumptions"
# in README.md.
ALIGNED = {
  ("Cusher",         "BERRY, Paul"):          "DUP",
  ("The Moor",       "Donnelly, Gary"):       "Sinn Féin",
  ("Oldpark",        "McCusker, Paul"):       "SDLP",
  ("Newtownards",    "IRVINE, Steven Gary"):  "DUP",
  ("Bangor Central", "IRVINE, Wesley Graham"): "DUP",
  ("Dungannon",      "MONTEITH, Barry"):      "Sinn Féin",
  ("Bann",           "McQUILLAN, Adrian"):    "DUP",
}

# Manual party assignments, keyed by (DEA, candidate). Used only where the
# source itself carries no description; never to override what a sheet states.
# Sperrin's workbook leaves the Description cell as 0 for these two; both
# confirmed Independent by the user (2026-09-10). Recorded per candidate as
# party_source="manual" so the assignment stays visible downstream.
MANUAL_PARTY = {
  ("Sperrin", "Barr, Raymond"):     "Independent",
  ("Sperrin", "Gallagher, Paul"):   "Independent",
}

def standardise(raw):
    """Map a source description to a party code, or None if unrecognised."""
    k = norm_ws(raw).lower().strip()
    if k in PARTY_MAP:
        return PARTY_MAP[k]
    # only then try without a trailing dot ("D.U.P" vs "D.U.P.")
    return PARTY_MAP.get(k.rstrip("."), PARTY_MAP.get(k + ".", None))

# ------------------------------------------------------------- BUILD -------
DEA_GEOJSON = "/home/matt/vote/data/DEA2021.geojson"

def dea_key(name):
    """Normalise a DEA name for matching ('Holywood & Clandeboye' == '... and ...')."""
    k = norm_ws(name or "").lower().replace("&", " and ")
    return re.sub(r"[^a-z]", "", k)

def canonical_deas():
    with open(DEA_GEOJSON) as fh:
        g = json.load(fh)
    names = [f["properties"]["FinalR_DEA"] for f in g["features"]]
    return {dea_key(n): n for n in names}

def file_stem(path):
    b = os.path.basename(path)
    b = re.sub(r"^local-council-elections-2023-result-sheets?-", "", b)
    return re.sub(r"(_\d+)?\.(pdf|xlsx)$", "", b)

def build():
    canon = canonical_deas()
    deas, flags = {}, []

    def add(dea_name, meta, rows, src, fmt):
        name = canon.get(dea_key(dea_name))
        if name is None:
            flags.append({"level": "error", "dea": dea_name, "source": src,
                          "issue": "DEA name does not match any DEA2021 boundary"})
            name = dea_name
        cands = []
        for r in rows:
            party = standardise(r["party_raw"])
            party_source = "source"
            if party is None:
                manual = MANUAL_PARTY.get((name, r["candidate"]))
                if manual is not None:
                    party, party_source = manual, "manual"
                    flags.append({"level": "info", "dea": name, "source": src,
                                  "issue": "party assigned manually; source gives none",
                                  "detail": f"{r['candidate']} -> {manual} "
                                            f"(sheet Description = {r['party_raw']!r})"})
                else:
                    flags.append({"level": "error", "dea": name, "source": src,
                                  "issue": "unrecognised party description",
                                  "detail": r["party_raw"], "candidate": r["candidate"],
                                  "first_pref": r["first_pref"]})
            row = {"candidate": r["candidate"], "party": party,
                   "party_raw": r["party_raw"], "party_source": party_source,
                   "first_pref": r["first_pref"]}
            aligned = ALIGNED.get((name, r["candidate"]))
            if aligned:
                row["party_as_declared"] = party
                row["party"], row["party_source"] = aligned, "assumed"
                flags.append({"level": "info", "dea": name, "source": src,
                              "issue": "party assigned by voting modelling assumption",
                              "detail": f"{r['candidate']}: {party} -> {aligned}"})
            cands.append(row)
        total = sum(c["first_pref"] for c in cands)
        valid = meta.get("valid_votes")
        ok = (valid is not None and total == valid)
        if not ok:
            flags.append({"level": "error", "dea": name, "source": src,
                          "issue": "first preferences do not sum to stated Total Valid Votes",
                          "detail": f"sum={total} stated={valid}"})
        parties = defaultdict(int)
        for c in cands:
            parties[c["party"] or "UNKNOWN"] += c["first_pref"]
        if name in deas:
            flags.append({"level": "error", "dea": name, "source": src,
                          "issue": "duplicate DEA"})
        deas[name] = {
            "dea": name, "source_file": src, "source_format": fmt,
            "seats": meta.get("seats"), "electorate": meta.get("eligible_electorate"),
            "votes_polled": meta.get("votes_polled"), "valid_votes": valid,
            "invalid_votes": meta.get("invalid_votes"), "quota": meta.get("quota"),
            "reconciled_to_stated_total": ok,
            "first_prefs_total": total,
            "party_first_prefs": dict(sorted(parties.items(), key=lambda kv: -kv[1])),
            "candidates": sorted(cands, key=lambda c: -c["first_pref"]),
        }

    for f in sorted(glob.glob(RAW + "/*.pdf")):
        try:
            dea, meta, rows = parse_pdf(f)
        except ScannedPDF:
            continue
        except Exception as e:
            flags.append({"level": "error", "source": os.path.basename(f),
                          "issue": "parse failed", "detail": str(e)})
            continue
        add(dea or file_stem(f), meta, rows, os.path.basename(f), "pdf_text")

    for f in sorted(glob.glob(RAW + "/*.xlsx")):
        try:
            sheets = parse_xlsx(f)
        except Exception as e:
            flags.append({"level": "error", "source": os.path.basename(f),
                          "issue": "parse failed", "detail": str(e)})
            continue
        fmt = "xlsx_bundle" if len(sheets) > 1 else "xlsx"
        for dea, meta, rows in sheets:
            rows = [r for r in rows
                    if not (r["candidate"] == "0" or
                            (r["party_raw"] == "0" and r["first_pref"] == 0))]
            add(dea or file_stem(f), meta, rows, os.path.basename(f), fmt)

    for dea, blob in SCANNED.items():
        rows = [{"candidate": c, "party_raw": p, "first_pref": v}
                for c, p, v in blob["rows"]]
        stem = dea.lower().replace(" ", "-")
        add(dea, blob["meta"], rows,
            f"local-council-elections-2023-result-sheet-{stem}.pdf",
            "pdf_scanned_transcribed")

    missing = [n for n in canonical_deas().values() if n not in deas]
    for m in missing:
        flags.append({"level": "error", "dea": m, "issue": "DEA missing from output"})

    party_totals = defaultdict(int)
    for d in deas.values():
        for p, v in d["party_first_prefs"].items():
            party_totals[p] += v

    out = {
        "election": "Northern Ireland local council elections 2023",
        "poll_date": "2023-05-18",
        "source": "https://www.eoni.org.uk/results-data/local-council-elections-2023-results/",
        "measure": "first preference votes by party, by district electoral area",
        "keyed_by": "DEA name, as FinalR_DEA in DEA2021.geojson",
        "notes": {
            "party": "Short label normalised from the sheet's own description, which is "
                     "kept verbatim in each candidate's `party_raw`. Labels and "
                     "`party_names` match data/pc24_results_2024.json exactly for all 12 "
                     "parties common to both, so the two elections join on `party`. "
                     "IRSP, PUP, Socialist Party and Workers Party stood at council "
                     "level only.",
            "party_source": "'source' where the party comes from the sheet's own "
                            "description; 'assumed' where a voting modelling assumption "
                            "counts the candidate as a different party, with the declared "
                            "label kept in `party_as_declared`; 'manual' where the sheet "
                            "gives no party at all and one was assigned by hand. Both "
                            "non-source cases are listed in `flags`. See \"Voting "
                            "modelling assumptions\" in README.md.",
            "verification": "Each DEA's first preferences are summed and compared with "
                            "the Total Valid Votes printed on the same sheet; all 80 "
                            "reconcile (`reconciled_to_stated_total`).",
            "candidate_names": "As printed on the result sheet, usually 'Surname, Forename'.",
        },
        "dea_count": len(deas),
        "candidate_count": sum(len(d["candidates"]) for d in deas.values()),
        "reconciled_dea_count": sum(1 for d in deas.values()
                                    if d["reconciled_to_stated_total"]),
        "party_names": PARTY_CANON,
        "ni_party_first_prefs": dict(sorted(party_totals.items(), key=lambda kv: -kv[1])),
        "flags": flags,
        "deas": [deas[k] for k in sorted(deas)],
    }
    with open(OUT, "w") as fh:
        json.dump(out, fh, indent=1, ensure_ascii=False)
    return out

if __name__ == "__main__":
    o = build()
    print(f"DEAs: {o['dea_count']}/80   candidates: {o['candidate_count']}")
    print(f"reconciled to stated Total Valid Votes: {o['reconciled_dea_count']}/{o['dea_count']}")
    print(f"flags: {len(o['flags'])}")
    for f in o["flags"]:
        print("  !", {k: v for k, v in f.items() if k != "level"})
    print("\nNI-wide first preferences:")
    tot = sum(o["ni_party_first_prefs"].values())
    for p, v in o["ni_party_first_prefs"].items():
        print(f"  {p:8} {v:>8}  {100*v/tot:5.2f}%   {PARTY_CANON.get(p,'?')}")
    print(f"  {'TOTAL':8} {tot:>8}")
