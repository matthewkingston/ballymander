#!/usr/bin/env python3
"""Extract STV transfer events from the 2022 Assembly and 2023 council count sheets.

First preferences remain the measure of voting intention everywhere else in this
project. Transfers are used for one thing only: estimating where a party's voters
go when that party isn't on the ballot (the transfer matrix, fit_transfer_matrix.py).

Each count sheet is reduced to an ordered list of stages -- every candidate's
transfer and the non-transferable votes -- and one routine turns stages into
events. An event is one stage with a single source candidate:

  source      the candidate whose votes moved (the one with a negative transfer)
  kind        "surplus" if the source had reached quota, else "exclusion". This
              agrees with the sheet's own Transfer/Exclude label wherever one exists.
  continuing  candidates still able to receive: not excluded, and below quota
  other_party_share
              how much of what was passed on came from other parties' voters.
              Exclusion: votes received from other-party sources / current total
              (first preferences and same-party transfers count as the party's
              own). Surplus: 0 if elected on first preferences, otherwise judged
              by the last parcel received -- in NI only that parcel is examined.

Stages are skipped, and counted, when several candidates were excluded together
(the flows can't be split) or when transfers plus non-transferable don't net to
zero (an undistributed final stage, or an inconsistent sheet).

Sources: all 18 Assembly workbooks; the 38 council DEAs published as xlsx; the 35
text-layer council PDFs, read with the column/row machinery of
parse_council_2023.py. The 7 Fermanagh & Omagh sheets are scanned images and are
not included. PDF columns are aligned by self-validation: first preferences must
match council_elections_2023.json, and every printed running total must equal the
previous total plus that stage's transfer, or the sheet is rejected.

Candidate parties, including the modelling assumptions for independents, come
from the existing first-preference JSONs.

Reads   data/assembly_elections_22/{assembly_2022.json, raw/*.xlsx}
        data/council_elections_23/{council_elections_2023.json, raw/*}
Writes  data/model/transfer_events.json
"""
from __future__ import annotations

import collections
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from parse_council_2023 import cell_text, column_rules, row_bands, row_cells  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ASSEMBLY = os.path.join(ROOT, "data", "assembly_elections_22")
COUNCIL = os.path.join(ROOT, "data", "council_elections_23")
OUT = os.path.join(ROOT, "data", "model", "transfer_events.json")
NT_LABEL = re.compile(r"non[- ]?transfer", re.I)


def key(s) -> str:
    return re.sub(r"[^a-z]", "", str(s).lower().replace("&", "and"))


def num(v):
    if v is None or v == "":
        return 0.0
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return None


# ------------------------------------------------------------------ xlsx ----
def xlsx_stages(ws, cand):
    """Stages from an EONI workbook sheet; long counts continue in a second block."""
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    stages, seen = [], set()
    for srow, r in enumerate(grid):
        if not any(v is not None and re.fullmatch(r"Stage\s+1", str(v).strip()) for v in r[:6]):
            continue
        cols = [(int(str(v).split()[1]), j) for j, v in enumerate(r[:23])
                if v is not None and re.fullmatch(r"Stage\s+\d+", str(v).strip())]
        rows, nt = {}, None
        for rr in grid[srow + 3:srow + 50]:
            if NT_LABEL.search(" ".join(str(v) for v in rr[:5] if v is not None)):
                nt = rr
                break
            for nc in (1, 2):                     # names in column C, or B (Belfast)
                if nc < len(rr) and rr[nc] is not None and key(rr[nc]) in cand:
                    rows[key(rr[nc])] = rr
                    break
        if len(rows) != len(cand) or nt is None:
            raise ValueError(f"matched {len(rows)}/{len(cand)} candidates, NT row {nt is not None}")
        for k, col in cols:
            if k == 1 or k in seen:           # the second block repeats the last stage
                continue
            seen.add(k)
            stages.append({"stage": k, "label": str(grid[srow + 1][col] or "").strip().lower(),
                           "tr": {n: num(rows[n][col]) or 0.0 for n in rows},
                           "nt": num(nt[col]) or 0.0})
    return stages


# ------------------------------------------------------------------- pdf ----
def pdf_page_rows(page):
    words = page.extract_words(x_tolerance=1)
    cand_w = [w for w in words if w["text"] == "Candidate"]
    if not cand_w:
        return None
    hdr = cand_w[0]["top"]
    rules = column_rules(page, hdr)
    party_w = [w for w in words if w["text"] in ("Description", "Party") and abs(w["top"] - hdr) < 4]

    def col_of(x):
        return max([j for j, r in enumerate(rules) if x >= r - 0.5] or [0])
    party_c = col_of(party_w[0]["x0"])
    bands = row_bands(page, hdr)
    by_band = collections.defaultdict(list)
    for ch in page.chars:
        mid = (ch["top"] + ch["bottom"]) / 2
        if mid > hdr + 3:
            by_band[max([j for j, y in enumerate(bands) if mid >= y - 0.5] or [-1])].append(ch)
    rows = []
    for b in sorted(by_band):
        cells = row_cells(by_band[b], rules)
        vals = {k: num(cell_text(cells[k])) for k in cells if k > party_c}
        rows.append({"label": " ".join(cell_text(cells[k]) for k in sorted(cells) if k <= party_c),
                     "vals": {k: v for k, v in vals.items() if v is not None}})
    return rows, party_c, len(rules)


def pdf_stages(path, cand):
    import pdfplumber
    stages = []
    total = {n: float(c["first_pref"]) for n, c in cand.items()}
    with pdfplumber.open(path) as pdf:
        for pno, page in enumerate(pdf.pages):
            got = pdf_page_rows(page)
            if not got:
                continue
            rows, party_c, ncol = got
            nts = [r for r in rows if NT_LABEL.search(r["label"])]
            if not nts:
                raise ValueError(f"page {pno + 1}: no non-transferable row")
            body = [r for r in rows if r["vals"] and not NT_LABEL.search(r["label"])
                    and not re.search(r"totals?", r["label"], re.I)]
            # candidates: by name; then by first preference (wrapped names); then the one left
            match, left = {}, set(cand)
            for i, r in enumerate(body):
                nm = key(re.sub(r"^(elected|excluded|e)\b|\d+", "", r["label"], flags=re.I))
                hit = [n for n in left if n and n in nm]
                if hit:
                    match[i] = max(hit, key=len)
                    left.discard(match[i])
            if pno == 0:
                for i, r in enumerate(body):
                    if i not in match:
                        hit = [n for n in left if any(abs(v - cand[n]["first_pref"]) < 1e-6
                                                      for v in r["vals"].values())]
                        if len(hit) == 1:
                            match[i] = hit[0]
                            left.discard(hit[0])
            unmatched = [i for i in range(len(body)) if i not in match]
            if len(unmatched) == 1 and len(left) == 1:
                match[unmatched.pop()] = left.pop()
            if unmatched or (left and pno == 0):
                raise ValueError(f"page {pno + 1}: unmatched rows {len(unmatched)}, candidates {sorted(left)}")
            grid = {match[i]: body[i]["vals"] for i in match}
            for n in left:                       # blank row on a continuation page
                grid[n] = {}
            # (transfer, total) column pairs start where the running totals chain
            best = None
            for start in range(party_c + 1, party_c + 4):
                ok = bad = 0
                run = dict(total)
                for tc in range(start, ncol, 2):
                    for n in grid:
                        run[n] += grid[n].get(tc, 0.0)
                        if tc + 1 in grid[n]:
                            ok, bad = (ok + 1, bad) if abs(grid[n][tc + 1] - run[n]) < 0.02 else (ok, bad + 1)
                if best is None or ok - 5 * bad > best[0]:
                    best = (ok - 5 * bad, start, bad)
            _, start, bad = best
            if bad:
                raise ValueError(f"page {pno + 1}: {bad} running totals disagree")
            for tc in range(start, ncol, 2):
                tr = {n: grid[n].get(tc, 0.0) for n in grid}
                if not any(abs(v) > 1e-9 for v in tr.values()):
                    continue
                for n in tr:
                    total[n] += tr[n]
                stages.append({"stage": len(stages) + 2, "label": "", "tr": tr,
                               "nt": nts[0]["vals"].get(tc, 0.0)})
    return stages


# ---------------------------------------------------------------- events ----
def events_from(election, area, meta, stages, log):
    cand = {key(x["candidate"]): x for x in meta["candidates"]}
    quota = meta.get("quota") or (meta["valid_votes"] // (meta["seats"] + 1) + 1)   # Droop

    def party(n):
        return cand[n]["party"]
    total = {n: float(x["first_pref"]) for n, x in cand.items()}
    other = {n: 0.0 for n in cand}
    parcel_other = {n: 0.0 for n in cand}
    excluded, events = set(), []
    for st in stages:
        tr, nt = st["tr"], st["nt"]
        out = [n for n in tr if tr[n] < -1e-6]
        if not out:
            continue
        moved = -sum(tr[n] for n in out)
        prev, prev_other = dict(total), dict(other)
        src_parties = {party(m) for m in out}
        for n in tr:
            total[n] = prev[n] + tr[n]
            own = src_parties == {party(n)}
            if tr[n] > 0 and not own:
                other[n] += tr[n]
            if prev[n] < quota - 1e-6 <= total[n]:
                parcel_other[n] = 0.0 if own else 1.0
        is_exclusion = all(prev[m] < quota - 1e-6 for m in out)
        if st["label"]:
            log["label agrees with quota rule"][st["label"].startswith("excl") == is_exclusion] += 1
        if is_exclusion:
            excluded.update(out)
        if abs(sum(tr.values()) + nt) > 0.01 * moved:
            log["skipped"][f"{election}: undistributed or unbalanced"] += 1
            continue
        if len(out) > 1:
            log["skipped"][f"{election}: several excluded together"] += 1
            continue
        src = out[0]
        if is_exclusion:
            other_share = prev_other[src] / prev[src]
        else:
            other_share = 0.0 if cand[src]["first_pref"] >= quota - 1e-6 else parcel_other[src]
        continuing = [n for n in tr if n != src and n not in excluded and prev[n] < quota - 1e-6]
        if abs(sum(v for n, v in tr.items() if n != src and n not in continuing)) > 0.01:
            log["skipped"][f"{election}: votes to non-continuing candidates"] += 1
            continue
        events.append({
            "election": election, "area": area, "stage": st["stage"],
            "kind": "exclusion" if is_exclusion else "surplus",
            "source": cand[src]["candidate"], "party": party(src),
            "votes_moved": round(moved, 2), "non_transferable": round(nt, 2),
            "other_party_share": round(other_share, 4),
            "continuing": [{"candidate": cand[n]["candidate"], "party": party(n),
                            "votes": round(tr[n], 2)} for n in continuing],
        })
    return events


def units():
    import openpyxl
    assembly = json.load(open(os.path.join(ASSEMBLY, "assembly_2022.json")))
    for c in assembly["constituencies"]:
        wb = openpyxl.load_workbook(os.path.join(ASSEMBLY, "raw", c["source_file"]), data_only=True)
        yield "assembly 2022", c["constituency"], c, lambda cand, ws=wb.worksheets[0]: xlsx_stages(ws, cand)
    council = json.load(open(os.path.join(COUNCIL, "council_elections_2023.json")))
    books = {}
    for d in council["deas"]:
        path = os.path.join(COUNCIL, "raw", d["source_file"])
        if d["source_format"] in ("xlsx", "xlsx_bundle"):
            wb = books.setdefault(path, openpyxl.load_workbook(path, data_only=True))
            sheets = wb.worksheets if len(wb.worksheets) == 1 else \
                [w for w in wb.worksheets if key(w.title) == key(d["dea"])]
            yield "council 2023", d["dea"], d, lambda cand, ws=sheets[0]: xlsx_stages(ws, cand)
        elif d["source_format"] == "pdf_text":
            yield "council 2023", d["dea"], d, lambda cand, p=path: pdf_stages(p, cand)


def main() -> None:
    log = collections.defaultdict(collections.Counter)
    events, areas = [], collections.Counter()
    for election, area, meta, read in units():
        cand = {key(x["candidate"]): x for x in meta["candidates"]}
        try:
            stages = read(cand)
        except ValueError as ex:
            raise SystemExit(f"{election} {area}: {ex}")
        events += events_from(election, area, meta, stages, log)
        areas[election] += 1
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w") as fh:
        json.dump({
            "description": "single-source STV transfer events (see scripts/parse_transfers.py)",
            "areas": dict(areas),
            "skipped_stages": dict(log["skipped"]),
            "events": events,
        }, fh, ensure_ascii=False)
        fh.write("\n")
    if log["label agrees with quota rule"][False]:
        raise SystemExit(f"surplus/exclusion rule disagrees with sheet labels "
                         f"{log['label agrees with quota rule'][False]} times")
    print(f"areas: {dict(areas)}")
    print(f"events: {dict(collections.Counter(e['election'] for e in events))}")
    print(f"skipped stages: {dict(log['skipped'])}")
    print(f"surplus/exclusion rule agrees with all {log['label agrees with quota rule'][True]} sheet labels")
    print(f"wrote {os.path.relpath(OUT, ROOT)}")


if __name__ == "__main__":
    main()
